#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
csv_to_excel_safe.py
====================
CSV dosyalarını Excel'in hücreleri YANLIŞLIKLA TARİHE çevirmesini önleyen
.xlsx dosyalarına dönüştürür.

Neden gerekli?
  CSV düz metindir ve hücre tipi bilgisi taşımaz; bu yüzden Excel dosyayı
  açarken "1.A.3", "3/7", "5-1" gibi değerleri TAHMİNEN tarihe çevirir ve
  kaydedince orijinal değer kalıcı bozulur. .xlsx ise her hücrenin tipini
  saklar: metin olarak yazılan kod hücresi tarihe çevrilmez.

Mantık:
  - Her kolon önce METİN olarak okunur (hiçbir değer bozulmaz).
  - TAMAMEN sayısal olan kolonlar sayıya çevrilir (analiz için).
  - Geri kalan (kod/etiket) kolonlar METİN olarak yazılır -> tarih sorunu biter.

Kurulum:
    pip install pandas openpyxl

Kullanım:
    python csv_to_excel_safe.py panel.csv
    python csv_to_excel_safe.py edgar_output/*.csv
    python csv_to_excel_safe.py veri.csv --text-cols ipcc2006_code sector substance
"""

from __future__ import annotations

import argparse
import glob
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

EXCEL_ROW_LIMIT = 1_048_576


def _coerce_numeric_columns(df: pd.DataFrame, force_text: set[str]) -> pd.DataFrame:
    """Tamamen sayısal kolonları sayıya çevirir; gerisi metin kalır."""
    for col in df.columns:
        if col in force_text:
            continue
        s = df[col]
        num = pd.to_numeric(s, errors="coerce")
        # Boş olmayan tüm değerler sayıya çevrilebiliyorsa kolon sayısaldır
        non_null = s.notna().sum()
        if non_null > 0 and num.notna().sum() == non_null:
            df[col] = num
    return df


def convert(csv_path: Path, force_text: set[str]) -> Path:
    # dtype=str: okuma sırasında hiçbir değer dönüştürülmez (kodlar korunur)
    df = pd.read_csv(csv_path, dtype=str, keep_default_na=False, na_values=[""])
    df = _coerce_numeric_columns(df, force_text)

    if len(df) > EXCEL_ROW_LIMIT:
        raise ValueError(f"{csv_path.name}: {len(df):,} satır Excel sınırını aşıyor "
                         f"(>{EXCEL_ROW_LIMIT:,}). Önce filtrele/böl.")

    out = csv_path.with_suffix(".xlsx")
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        df.to_excel(xw, index=False, sheet_name="data")
        ws = xw.sheets["data"]
        # Metin kolonlarının hücre biçimini açıkça '@' (metin) yap -> ekstra garanti
        text_cols = [c for c in df.columns
                     if c in force_text or df[c].dtype == object]
        for col in text_cols:
            j = df.columns.get_loc(col) + 1
            letter = get_column_letter(j)
            for row in range(2, len(df) + 2):           # başlık 1. satır
                ws[f"{letter}{row}"].number_format = "@"
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description="CSV -> Excel-güvenli .xlsx (tarih bozulması yok).")
    ap.add_argument("paths", nargs="+", help="CSV dosya(lar)ı veya joker (ör. *.csv)")
    ap.add_argument("--text-cols", nargs="*", default=[],
                    help="Her zaman METİN kalsın denen kolon adları "
                         "(ör. ipcc2006_code sector substance iso3)")
    args = ap.parse_args()

    # Joker desenleri genişlet
    files: list[Path] = []
    for p in args.paths:
        files.extend(Path(f) for f in glob.glob(p))
    if not files:
        print("Eşleşen CSV bulunamadı.")
        return

    force_text = set(args.text_cols)
    for f in files:
        try:
            out = convert(f, force_text)
            print(f"[OK] {f}  ->  {out}")
        except Exception as e:
            print(f"[HATA] {f}: {e}")


if __name__ == "__main__":
    main()
