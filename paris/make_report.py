#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
make_report.py
==============
analysis_pipeline_full.py tahmincilerini MAKALEYE HAZIR tablolara döker:
  * Tahminci karşılaştırma tablosu (Pooled / FE / RE / Anderson-Hsiao / MG / CCEMG)
    — katsayı + yıldızlı anlamlılık + parantezde standart hata
  * Tanı testleri özeti (Pesaran CD, Hausman, panel birim kök)

Çıktılar (analysis_output/):
  results_table.csv        — birleşik tablo (ham)
  regression_table.tex     — derlenebilir LaTeX (booktabs + threeparttable)
  regression_table.docx    — biçimlendirilmiş Word tablosu

Kurulum:
    pip install pandas numpy scipy statsmodels linearmodels python-docx
Kullanım:
    python make_report.py panel_output/panel_clean.csv
    python make_report.py panel_clean.csv --dep ln_co2_pc
"""

from __future__ import annotations
import argparse
from pathlib import Path
import numpy as np
import pandas as pd

import analysis_pipeline_full as A   # tahminci fonksiyonlarını yeniden kullan

REG_LABELS = {
    "ln_gdp_pc": "ln(GSYİH pc)", "ln_gdp_pc_sq": "ln(GSYİH pc)²",
    "ln_energy_pc": "ln(Enerji pc)", "energy_intensity": "Enerji yoğunluğu",
    "renew_wb_pct": "Yenilenebilir (%)", "urban_pct": "Kentleşme (%)",
    "trade_openness": "Ticari açıklık", "fdi_pct_gdp": "DYY (% GSYİH)",
    "fin_dev_credit": "Finansal gelişme", "industry_pct_gdp": "Sanayi (% GSYİH)",
}
LAG_LABEL = "Gecikmeli bağımlı (ρ)"


def stars(p):
    return "***" if p < 0.01 else "**" if p < 0.05 else "*" if p < 0.10 else ""


# --------------------------------------------------------------------------- #
# Her tahminciden (katsayı, se, p) çıkar
# --------------------------------------------------------------------------- #
def from_linearmodels(res, regs):
    return {r: (res.params[r], res.std_errors[r], res.pvalues[r])
            for r in regs if r in res.params.index}


def from_ah(res, regs):
    out = {}
    for r in regs:
        k = "d_" + r
        if k in res.params.index:
            out[r] = (res.params[k], res.std_errors[k], res.pvalues[k])
    if "dy_lag1" in res.params.index:
        out[LAG_LABEL] = (res.params["dy_lag1"], res.std_errors["dy_lag1"],
                          res.pvalues["dy_lag1"])
    return out


def from_mg(tbl, regs):
    return {r: (tbl.loc[r, "coef"], tbl.loc[r, "se"], tbl.loc[r, "p"])
            for r in regs if r in tbl.index}


# --------------------------------------------------------------------------- #
# Birleşik tablo kur
# --------------------------------------------------------------------------- #
def build_results(frame, dep, regs):
    pooled, fe_cl, fe_dk, re, hau = A.panel_models(frame, dep, regs)
    ah = A.anderson_hsiao(frame, dep, regs)
    mg, cce = A.mg_ccemg(frame, dep, regs)

    models = {
        "Havuzlanmış OLS": (from_linearmodels(pooled, regs), pooled.nobs,
                            frame["iso3"].nunique(), pooled.rsquared),
        "Sabit Etki (FE)": (from_linearmodels(fe_cl, regs), fe_cl.nobs,
                            frame["iso3"].nunique(), fe_cl.rsquared_within),
        "Rassal Etki (RE)": (from_linearmodels(re, regs), re.nobs,
                             frame["iso3"].nunique(), re.rsquared),
        "Anderson-Hsiao": (from_ah(ah, regs), int(ah.nobs),
                           frame["iso3"].nunique(), np.nan),
        "Mean Group (MG)": (from_mg(mg, regs), len(frame),
                            int(mg["N"].iloc[0]), np.nan),
        "CCEMG": (from_mg(cce, regs), len(frame),
                  int(cce["N"].iloc[0]), np.nan),
    }

    row_order = regs + [LAG_LABEL]
    coef_rows, se_rows = [], []
    for var in row_order:
        cr = {"variable": REG_LABELS.get(var, var)}
        sr = {"variable": ""}
        for mname, (d, *_ ) in models.items():
            if var in d:
                c, s, p = d[var]
                cr[mname] = f"{c:.3f}{stars(p)}"
                sr[mname] = f"({s:.3f})"
            else:
                cr[mname] = ""; sr[mname] = ""
        coef_rows.append(cr); se_rows.append(sr)

    # alt bilgi satırları
    foot = []
    for label, key in [("Gözlem (N)", 1), ("Ülke", 2), ("R² / within", 3)]:
        row = {"variable": label}
        for mname, tup in models.items():
            val = tup[key]
            row[mname] = ("" if (isinstance(val, float) and np.isnan(val))
                          else (f"{val:.3f}" if key == 3 else f"{int(val):,}"))
        foot.append(row)

    # satırları sırala: her değişken için coef + se
    inter = []
    for cr, sr in zip(coef_rows, se_rows):
        if any(cr[m] for m in models):      # boş değişkenleri atla
            inter.append(cr); inter.append(sr)
    table = pd.DataFrame(inter + foot)
    cols = ["variable"] + list(models.keys())
    return table[cols], hau, list(models.keys())


def build_diagnostics(frame, dep, regs, hau):
    cd = A.pesaran_cd(frame, dep)
    ur_l = A.panel_unit_root(frame, dep)
    ur_d = A.panel_unit_root(frame, dep, diff=True)
    rows = [
        {"Test": "Pesaran CD (yatay kesit bağımlılığı)",
         "İstatistik": f"{cd['CD']:.2f}", "p": f"{cd['p_value']:.3f}",
         "Sonuç": "Bağımlılık var" if cd['p_value'] < 0.05 else "Yok"},
        {"Test": "Hausman (FE vs RE)",
         "İstatistik": f"{hau['chi2']:.2f}", "p": f"{hau['p_value']:.3f}",
         "Sonuç": "FE tercih" if hau['p_value'] < 0.05 else "RE tercih"},
        {"Test": f"Panel birim kök — {dep} (düzey, Fisher)",
         "İstatistik": f"χ²={ur_l['fisher_chi2']:.1f}", "p": f"{ur_l['p_value']:.3f}",
         "Sonuç": "Durağan" if ur_l['p_value'] < 0.05 else "Birim kök"},
        {"Test": f"Panel birim kök — {dep} (1. fark)",
         "İstatistik": f"χ²={ur_d['fisher_chi2']:.1f}", "p": f"{ur_d['p_value']:.3f}",
         "Sonuç": "Durağan" if ur_d['p_value'] < 0.05 else "Birim kök"},
    ]
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------- #
# LaTeX
# --------------------------------------------------------------------------- #
def write_latex(table, diag, dep, models, path):
    ncol = len(models)
    col_fmt = "l" + "c" * ncol
    esc = lambda s: str(s).replace("_", r"\_")    # LaTeX alt çizgi kaçışı
    head = " & ".join(["Değişken"] + models) + r" \\"
    body = []
    for _, r in table.iterrows():
        cells = [esc(r["variable"])] + [esc(r[m]) for m in models]
        body.append(" & ".join(cells) + r" \\")
    # alt bilgi ayıracı: son 3 satır öncesine \midrule koy
    body.insert(len(body) - 3, r"\midrule")

    dhead = r"Test & İstatistik & $p$ & Sonuç \\"
    dbody = [" & ".join([esc(r["Test"]), esc(r["İstatistik"]),
                         esc(r["p"]), esc(r["Sonuç"])]) + r" \\"
             for _, r in diag.iterrows()]

    tex = r"""\documentclass[11pt]{article}
\usepackage[utf8]{inputenc}
\usepackage[T1]{fontenc}
\usepackage{booktabs}
\usepackage{threeparttable}
\usepackage[margin=2cm,landscape]{geometry}
\renewcommand{\arraystretch}{1.05}
\begin{document}

\begin{table}[t]\centering
\begin{threeparttable}
\caption{Bağımlı değişken: %s --- Tahminci karşılaştırması}
\begin{tabular}{%s}
\toprule
%s
\midrule
%s
\bottomrule
\end{tabular}
\begin{tablenotes}\footnotesize
\item Parantez içinde standart hatalar. $^{***}p<0.01$, $^{**}p<0.05$, $^{*}p<0.10$.
\item Pooled/FE/RE: ülke düzeyinde kümelenmiş SE. FE ayrıca Driscoll-Kraay SE ile
de tahmin edilmiştir (yatay kesit bağımlılığına dayanıklı). Anderson-Hsiao: birinci
farklarda IV (kısa dönem). CCEMG: yatay kesit bağımlılığına dayanıklı; CD anlamlı
olduğunda esas alınır. Gelir terimleri ortalamadan arındırılmıştır (mean-centered).
\end{tablenotes}
\end{threeparttable}
\end{table}

\begin{table}[t]\centering
\caption{Tanı testleri}
\begin{tabular}{lccl}
\toprule
%s
\midrule
%s
\bottomrule
\end{tabular}
\end{table}

\end{document}
""" % (dep.replace("_", r"\_"), col_fmt, head, "\n".join(body),
       dhead, "\n".join(dbody))
    # pdflatex Yunan harflerini metin modunda desteklemez -> math moduna çevir
    tex = tex.replace("ρ", r"$\rho$").replace("χ²", r"$\chi^2$")
    Path(path).write_text(tex, encoding="utf-8")


# --------------------------------------------------------------------------- #
# Word (python-docx)
# --------------------------------------------------------------------------- #
def write_docx(table, diag, dep, models, path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    def shade(cell, fill="D9E2F3"):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear")
        sh.set(qn("w:fill"), fill); tcPr.append(sh)

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    doc.add_heading(f"Tablo 1. Bağımlı değişken: {dep} — Tahminci karşılaştırması", level=2)
    t = doc.add_table(rows=1, cols=len(models) + 1)
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    hdr[0].text = "Değişken"
    for j, m in enumerate(models):
        hdr[j + 1].text = m
    for c in hdr:
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True; run.font.size = Pt(9)
        shade(c)
    for _, r in table.iterrows():
        cells = t.add_row().cells
        cells[0].text = str(r["variable"])
        for j, m in enumerate(models):
            cells[j + 1].text = str(r[m])
            for p in cells[j + 1].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9)
    note = doc.add_paragraph()
    note.add_run("Not: ").bold = True
    note.add_run("Parantez içinde standart hatalar. ***p<0.01, **p<0.05, *p<0.10. "
                 "Pooled/FE/RE ülke düzeyinde kümelenmiş SE; FE ayrıca Driscoll-Kraay "
                 "ile tahmin edilmiştir. Anderson-Hsiao birinci farklarda IV (kısa dönem). "
                 "CCEMG yatay kesit bağımlılığına dayanıklıdır. Gelir terimleri mean-centered.")
    for run in note.runs:
        run.font.size = Pt(8)

    doc.add_heading("Tablo 2. Tanı testleri", level=2)
    d = doc.add_table(rows=1, cols=4); d.style = "Table Grid"
    for j, h in enumerate(["Test", "İstatistik", "p", "Sonuç"]):
        d.rows[0].cells[j].text = h
        for p in d.rows[0].cells[j].paragraphs:
            for run in p.runs:
                run.bold = True; run.font.size = Pt(9)
        shade(d.rows[0].cells[j])
    for _, r in diag.iterrows():
        cells = d.add_row().cells
        for j, k in enumerate(["Test", "İstatistik", "p", "Sonuç"]):
            cells[j].text = str(r[k])
            for p in cells[j].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.save(path)


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("csv")
    ap.add_argument("--dep", default="ln_co2_pc")
    a = ap.parse_args()

    out = Path(a.csv).parent / "analysis_output"; out.mkdir(exist_ok=True)
    frame, regs = A.prep_frame(pd.read_csv(a.csv), a.dep)
    print(f"Model: {a.dep} ~ {regs} | {frame['iso3'].nunique()} ülke, {len(frame):,} gözlem")

    table, hau, models = build_results(frame, a.dep, regs)
    diag = build_diagnostics(frame, a.dep, regs, hau)

    table.to_csv(out / "results_table.csv", index=False)
    write_latex(table, diag, a.dep, models, out / "regression_table.tex")
    write_docx(table, diag, a.dep, models, out / "regression_table.docx")
    print("Üretildi:", out / "results_table.csv", "|",
          out / "regression_table.tex", "|", out / "regression_table.docx")


if __name__ == "__main__":
    main()
